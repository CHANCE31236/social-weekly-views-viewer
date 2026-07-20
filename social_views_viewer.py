import csv
import json
import os
import re
import subprocess
import sys
import time
import traceback
from datetime import date, timedelta
from pathlib import Path
from tkinter import BOTH, LEFT, RIGHT, Menu, StringVar, Tk, Toplevel, messagebox
from tkinter import ttk
from urllib.parse import parse_qs, unquote, urlparse


APP_NAME = "Social Weekly Views Viewer"


LANGUAGES = {
    "en": "English",
    "zh": "中文",
    "fr": "Français",
}


TEXTS = {
    "en": {
        "title": APP_NAME,
        "advanced": "Advanced",
        "language": "Language",
        "date_range": "Date range",
        "reset_dates": "Reset last 7 days",
        "open_all_dashboards": "Open all dashboards",
        "open_all_public": "Open all public pages",
        "platform": "Platform",
        "account": "Account name",
        "views": "Views this week",
        "action": "Action",
        "open_dashboard": "Open dashboard",
        "open_public": "Public view",
        "read_views": "Read data",
        "export_report": "Export report",
        "add_account": "Add account",
        "edit_account": "Edit account",
        "delete_account": "Delete account",
        "clear_values": "Clear values",
        "open_config": "Open app folder",
        "ready": "Ready. Open dashboards, read data, then export report.",
        "date_reset": "Date range reset.",
        "dashboard_opened": "{platform} dashboard opened. Log in on the official site, make Views visible, then click Read data.",
        "public_missing": "{platform} has no public URL.",
        "public_opened": "{platform} public page opened. Make public views visible, then click Read data.",
        "all_dashboards_opened": "All dashboards opened. Log in and make weekly Views visible.",
        "all_public_opened": "All public pages opened.",
        "read_prompt_title": "Read data",
        "read_prompt": "Please make sure the {platform} dashboard or public page is open and the Views number is visible on screen.\n\nAfter clicking OK, switch to that browser page within 5 seconds.\n\nThis app never asks for or saves passwords.",
        "reading_wait": "Reading {platform} in 5 seconds. Bring the browser page to front...",
        "reading": "Reading {platform}...",
        "read_failed": "Could not read automatically. Please type the view count manually.",
        "detected": "{platform} Views detected: {views}",
        "exported": "Report exported.",
        "export_title": "Export report",
        "export_success": "Report exported to Desktop:\n{path}",
        "export_failed": "Export failed. Check last_error.log.",
        "clear_title": "Clear values",
        "clear_confirm": "Clear all weekly views values?",
        "cleared": "Values cleared.",
        "account_settings": "Account links",
        "dashboard_url": "Official dashboard / login URL",
        "public_url": "Public profile URL",
        "auto_detect": "Platform and account name will be detected automatically from the links.",
        "missing_url": "Please enter at least one URL.",
        "save": "Save",
        "cancel": "Cancel",
        "col_platform": "Platform",
        "col_account": "Account name",
        "col_views": "Views this week",
        "col_date": "Date range",
        "col_raw": "Raw input",
    },
    "zh": {
        "title": APP_NAME,
        "advanced": "高级",
        "language": "语言",
        "date_range": "日期范围",
        "reset_dates": "重置为最近 7 天",
        "open_all_dashboards": "打开全部后台",
        "open_all_public": "打开全部公开页",
        "platform": "平台",
        "account": "账号名称",
        "views": "本周浏览量",
        "action": "操作",
        "open_dashboard": "打开后台",
        "open_public": "公开页",
        "read_views": "读取数据",
        "export_report": "导出报表",
        "add_account": "添加账号",
        "edit_account": "编辑账号",
        "delete_account": "删除账号",
        "clear_values": "清空数字",
        "open_config": "打开程序文件夹",
        "ready": "准备好了。打开后台，读取数据，然后导出报表。",
        "date_reset": "日期范围已重置。",
        "dashboard_opened": "{platform} 后台已打开。请在官方网页登录并确认浏览量可见，然后点击读取数据。",
        "public_missing": "{platform} 没有公开页链接。",
        "public_opened": "{platform} 公开页已打开。请确认公开浏览量可见，然后点击读取数据。",
        "all_dashboards_opened": "全部后台已打开。请登录并确认最近 7 天浏览量可见。",
        "all_public_opened": "全部公开页已打开。",
        "read_prompt_title": "读取数据",
        "read_prompt": "请确认 {platform} 后台页面或公开页面已经打开，并且浏览量数字在屏幕上可见。\n\n点击 OK 后，请在 5 秒内切换到对应浏览器页面。\n\n本程序不会要求或保存密码。",
        "reading_wait": "将在 5 秒后读取 {platform}。请把浏览器页面切到最前面...",
        "reading": "正在读取 {platform}...",
        "read_failed": "未能自动识别，请手动输入浏览量。",
        "detected": "{platform} 浏览量已识别：{views}",
        "exported": "报表已导出。",
        "export_title": "导出报表",
        "export_success": "报表已导出到桌面：\n{path}",
        "export_failed": "导出失败，请检查 last_error.log。",
        "clear_title": "清空数字",
        "clear_confirm": "确定清空所有本周浏览量吗？",
        "cleared": "数字已清空。",
        "account_settings": "账号链接",
        "dashboard_url": "官方后台 / 登录链接",
        "public_url": "公开主页链接",
        "auto_detect": "平台和账号名称会根据链接自动识别。",
        "missing_url": "请至少输入一个链接。",
        "save": "保存",
        "cancel": "取消",
        "col_platform": "平台",
        "col_account": "账号名称",
        "col_views": "本周浏览量",
        "col_date": "日期范围",
        "col_raw": "原始输入",
    },
    "fr": {
        "title": APP_NAME,
        "advanced": "Avancé",
        "language": "Langue",
        "date_range": "Période",
        "reset_dates": "Réinitialiser 7 jours",
        "open_all_dashboards": "Ouvrir tous les tableaux",
        "open_all_public": "Ouvrir les pages publiques",
        "platform": "Plateforme",
        "account": "Nom du compte",
        "views": "Vues cette semaine",
        "action": "Action",
        "open_dashboard": "Ouvrir le tableau",
        "open_public": "Vue publique",
        "read_views": "Lire les données",
        "export_report": "Exporter le rapport",
        "add_account": "Ajouter un compte",
        "edit_account": "Modifier",
        "delete_account": "Supprimer",
        "clear_values": "Effacer les valeurs",
        "open_config": "Ouvrir le dossier",
        "ready": "Prêt. Ouvrez les tableaux, lisez les données, puis exportez le rapport.",
        "date_reset": "Période réinitialisée.",
        "dashboard_opened": "Tableau {platform} ouvert. Connectez-vous sur le site officiel, affichez les vues, puis cliquez sur Lire les données.",
        "public_missing": "{platform} n'a pas d'URL publique.",
        "public_opened": "Page publique {platform} ouverte. Affichez les vues publiques, puis cliquez sur Lire les données.",
        "all_dashboards_opened": "Tous les tableaux sont ouverts. Connectez-vous et affichez les vues des 7 derniers jours.",
        "all_public_opened": "Toutes les pages publiques sont ouvertes.",
        "read_prompt_title": "Lire les données",
        "read_prompt": "Vérifiez que le tableau ou la page publique {platform} est ouvert et que le nombre de vues est visible à l'écran.\n\nAprès OK, passez sur cette page du navigateur dans les 5 secondes.\n\nCette application ne demande et n'enregistre jamais les mots de passe.",
        "reading_wait": "Lecture de {platform} dans 5 secondes. Mettez la page du navigateur au premier plan...",
        "reading": "Lecture de {platform}...",
        "read_failed": "Lecture automatique impossible. Saisissez le nombre de vues manuellement.",
        "detected": "Vues {platform} détectées : {views}",
        "exported": "Rapport exporté.",
        "export_title": "Exporter le rapport",
        "export_success": "Rapport exporté sur le Bureau :\n{path}",
        "export_failed": "Échec de l'export. Vérifiez last_error.log.",
        "clear_title": "Effacer les valeurs",
        "clear_confirm": "Effacer toutes les vues hebdomadaires ?",
        "cleared": "Valeurs effacées.",
        "account_settings": "Liens du compte",
        "dashboard_url": "URL officielle du tableau / connexion",
        "public_url": "URL du profil public",
        "auto_detect": "La plateforme et le nom du compte seront détectés automatiquement depuis les liens.",
        "missing_url": "Veuillez saisir au moins une URL.",
        "save": "Enregistrer",
        "cancel": "Annuler",
        "col_platform": "Plateforme",
        "col_account": "Nom du compte",
        "col_views": "Vues cette semaine",
        "col_date": "Période",
        "col_raw": "Saisie brute",
    },
}


def app_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


APP_DIR = app_dir()
DATA_DIR = APP_DIR / "data"
EXPORT_DIR = APP_DIR / "exports"
DESKTOP = Path.home() / "Desktop"
CONFIG_FILE = APP_DIR / "accounts.json"
STATE_FILE = DATA_DIR / "weekly_views_state_v2.json"
ERROR_FILE = APP_DIR / "last_error.log"
SCREENSHOT_FILE = DATA_DIR / "last_autoread_screen.png"


DEFAULT_ACCOUNTS = [
    {
        "platform": "YouTube",
        "account": "Example account",
        "dashboard_url": "https://studio.youtube.com/",
        "public_url": "https://www.youtube.com/",
    },
    {
        "platform": "TikTok",
        "account": "Example account",
        "dashboard_url": "https://www.tiktok.com/tiktokstudio",
        "public_url": "https://www.tiktok.com/",
    },
    {
        "platform": "Instagram",
        "account": "Example account",
        "dashboard_url": "https://business.facebook.com/latest/insights",
        "public_url": "https://www.instagram.com/",
    },
]


def log_error(title: str = "") -> None:
    try:
        ERROR_FILE.write_text(
            (title + "\n\n" if title else "") + traceback.format_exc(),
            encoding="utf-8",
        )
    except Exception:
        pass


def ensure_dirs() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)


def load_accounts() -> list[dict]:
    ensure_dirs()
    if not CONFIG_FILE.exists():
        save_accounts(DEFAULT_ACCOUNTS)
        return list(DEFAULT_ACCOUNTS)
    try:
        data = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
        accounts = data.get("accounts", data if isinstance(data, list) else [])
        if not isinstance(accounts, list):
            raise ValueError("accounts.json must be a list or contain an accounts list")
        clean = []
        for item in accounts:
            if not isinstance(item, dict):
                continue
            dashboard_url = str(item.get("dashboard_url", "")).strip()
            public_url = str(item.get("public_url", item.get("guest_url", ""))).strip()
            inferred = infer_account(dashboard_url, public_url)
            clean.append(
                {
                    "platform": str(item.get("platform", "")).strip() or inferred["platform"],
                    "account": str(item.get("account", "")).strip() or inferred["account"],
                    "dashboard_url": dashboard_url,
                    "public_url": public_url,
                }
            )
        return clean or list(DEFAULT_ACCOUNTS)
    except Exception:
        log_error("Failed to load accounts.json")
        return list(DEFAULT_ACCOUNTS)


def save_accounts(accounts: list[dict]) -> None:
    ensure_dirs()
    CONFIG_FILE.write_text(
        json.dumps({"accounts": accounts}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def compact_account_name(value: str) -> str:
    value = unquote(value or "").strip().strip("/")
    value = value.lstrip("@")
    value = re.sub(r"[_\-]+", " ", value)
    value = re.sub(r"\s+", " ", value).strip()
    if not value:
        return ""
    if value.lower() in {"reel", "p", "shorts", "videos", "accounts", "analytics", "insights"}:
        return ""
    return value


def detect_platform(*urls: str) -> str:
    joined = " ".join(urls).lower()
    if "youtube.com" in joined or "youtu.be" in joined:
        return "YouTube"
    if "tiktok.com" in joined:
        return "TikTok"
    if "instagram.com" in joined or "business.facebook.com" in joined and "instagram" in joined:
        return "Instagram"
    if "facebook.com" in joined or "business.facebook.com" in joined:
        return "Facebook"
    if "linkedin.com" in joined:
        return "LinkedIn"
    if "x.com" in joined or "twitter.com" in joined:
        return "X"
    return "Other"


def detect_account_from_url(url: str) -> str:
    if not url:
        return ""
    try:
        parsed = urlparse(url)
    except Exception:
        return ""
    host = parsed.netloc.lower()
    path_parts = [part for part in parsed.path.split("/") if part]
    query = parse_qs(parsed.query)

    if "tiktok.com" in host:
        for part in path_parts:
            if part.startswith("@"):
                return compact_account_name(part)
    if "instagram.com" in host:
        if path_parts:
            return compact_account_name(path_parts[0])
    if "youtube.com" in host:
        for part in path_parts:
            if part.startswith("@"):
                return compact_account_name(part)
        if "channel" in path_parts:
            idx = path_parts.index("channel")
            if idx + 1 < len(path_parts):
                return compact_account_name(path_parts[idx + 1])
    if "facebook.com" in host:
        for key in ("asset_id", "page_id", "id"):
            if query.get(key):
                return compact_account_name(query[key][0])
        for part in path_parts:
            if part not in {"people", "latest", "insights", "results"}:
                return compact_account_name(part)
    for key in ("account", "user", "username", "handle"):
        if query.get(key):
            return compact_account_name(query[key][0])
    return compact_account_name(path_parts[0] if path_parts else host.replace("www.", ""))


def infer_account(dashboard_url: str, public_url: str) -> dict:
    platform = detect_platform(public_url, dashboard_url)
    account = detect_account_from_url(public_url) or detect_account_from_url(dashboard_url)
    if not account:
        account = f"{platform} account"
    return {
        "platform": platform,
        "account": account,
        "dashboard_url": dashboard_url.strip(),
        "public_url": public_url.strip(),
    }


def last_7_day_range() -> tuple[date, date]:
    end = date.today() - timedelta(days=1)
    start = end - timedelta(days=6)
    return start, end


def default_date_range() -> str:
    start, end = last_7_day_range()
    return f"{start.isoformat()} to {end.isoformat()}"


def parse_date_range(value: str) -> tuple[str, str]:
    matches = re.findall(r"\d{4}-\d{2}-\d{2}", value or "")
    if len(matches) >= 2:
        return matches[0], matches[1]
    start, end = last_7_day_range()
    return start.isoformat(), end.isoformat()


def load_state() -> dict:
    ensure_dirs()
    if not STATE_FILE.exists():
        return {"date_range": default_date_range(), "rows": {}}
    try:
        data = json.loads(STATE_FILE.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("state must be a dict")
        data.setdefault("date_range", default_date_range())
        data.setdefault("rows", {})
        return data
    except Exception:
        log_error("Failed to load state")
        return {"date_range": default_date_range(), "rows": {}}


def save_state(state: dict) -> None:
    ensure_dirs()
    STATE_FILE.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def edge_executable() -> str:
    candidates = [
        Path(os.environ.get("ProgramFiles(x86)", "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
        Path(os.environ.get("ProgramFiles", "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return "msedge"


def open_url(url: str) -> None:
    if not url:
        return
    try:
        subprocess.Popen(
            [edge_executable(), url],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except Exception:
        log_error("Failed to open URL")


def copy_active_window_text() -> str:
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-Command", "$wshell=New-Object -ComObject wscript.shell; $wshell.SendKeys('^a'); Start-Sleep -Milliseconds 100; $wshell.SendKeys('^c')"],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=3,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except Exception:
        log_error("Failed to copy active window text")
    return ""


def read_screen_ocr_items() -> list[dict]:
    try:
        from PIL import ImageGrab

        image = ImageGrab.grab(all_screens=True)
        image.save(SCREENSHOT_FILE)
        import easyocr

        reader = easyocr.Reader(["en"], gpu=False, verbose=False)
        result = reader.readtext(str(SCREENSHOT_FILE), detail=1, paragraph=False)
        items = []
        for box, text, confidence in result:
            xs = [point[0] for point in box]
            ys = [point[1] for point in box]
            items.append(
                {
                    "text": text,
                    "confidence": confidence,
                    "x": sum(xs) / len(xs),
                    "y": sum(ys) / len(ys),
                    "left": min(xs),
                    "right": max(xs),
                    "top": min(ys),
                    "bottom": max(ys),
                }
            )
        return items
    except Exception:
        log_error("OCR read failed")
        return []


def parse_number_token(text: str) -> tuple[str, int] | None:
    raw = (text or "").strip()
    if not raw:
        return None
    compact = raw.replace(",", "").replace(" ", "")
    if "+" in compact:
        total = 0
        for part in compact.split("+"):
            parsed = parse_number_token(part)
            if parsed is None:
                return None
            total += parsed[1]
        return raw, total
    match = re.search(r"(\d+(?:\.\d+)?)([kKmMwW万]?)", compact)
    if not match:
        return None
    value = float(match.group(1))
    suffix = match.group(2).lower()
    if suffix == "k":
        value *= 1000
    elif suffix == "m":
        value *= 1000000
    elif suffix in {"w", "万"}:
        value *= 10000
    return raw, int(round(value))


def safe_int(text: str) -> int:
    parsed = parse_number_token(text)
    return parsed[1] if parsed else 0


VIEW_KEYWORDS = [
    "views",
    "view",
    "video views",
    "post views",
    "reel views",
    "vues",
    "vue",
    "visualisations",
    "浏览",
    "浏览量",
    "观看",
    "播放",
    "播放量",
]

PRIMARY_VIEW_LABELS = [
    "video views",
    "videos views",
    "views",
    "vues vidéo",
    "vues des vidéos",
    "播放量",
    "视频播放量",
]

EXCLUDED_VIEW_LABELS = [
    "profile views",
    "profile view",
    "profile_view",
    "profile",
    "followers",
    "likes",
    "comments",
    "shares",
    "traffic source",
    "search queries",
]


def is_excluded_view_label(text: str) -> bool:
    low = (text or "").lower()
    return any(label in low for label in EXCLUDED_VIEW_LABELS)


def view_label_score(text: str) -> int | None:
    low = (text or "").lower()
    if is_excluded_view_label(low):
        return None
    if "video views" in low or "videos views" in low:
        return 100
    if "vues vidéo" in low or "vues des vidéos" in low:
        return 95
    if "视频播放量" in low:
        return 95
    if "reel views" in low or "post views" in low:
        return 80
    if "播放量" in low or "浏览量" in low:
        return 75
    if "views" in low or "vues" in low or "visualisations" in low:
        return 60
    if "view" in low or "vue" in low or "浏览" in low or "播放" in low:
        return 45
    return None


def parse_views_from_text(text: str) -> int | None:
    if not text:
        return None
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    best = None
    best_score = None
    for index, line in enumerate(lines):
        score = view_label_score(line)
        if score is None:
            continue
        nearby = " ".join(lines[index : index + 5])
        # Skip blocks where a higher-risk metric label appears before any number.
        if is_excluded_view_label(nearby) and "video views" not in nearby.lower() and "videos views" not in nearby.lower():
            continue
        for token in re.findall(r"\d+(?:[,\s]\d{3})*(?:\.\d+)?\s*[kKmMwW万]?", nearby):
            parsed = parse_number_token(token)
            if parsed:
                candidate_score = score * 1000000 - index
                if best_score is None or candidate_score > best_score:
                    best_score = candidate_score
                    best = parsed[1]
                break
    return best


def combined_number_value(item: dict, items: list[dict]) -> int | None:
    parsed = parse_number_token(item["text"])
    value = parsed[1] if parsed else None
    # EasyOCR sometimes splits "1.8K" into "1." and "8K". Recombine close tokens.
    same_line = [
        other
        for other in items
        if other is not item and abs(other.get("y", 0) - item.get("y", 0)) < 18
    ]
    left_neighbors = [
        other
        for other in same_line
        if 0 <= item.get("left", 0) - other.get("right", 0) <= 24
    ]
    right_neighbors = [
        other
        for other in same_line
        if 0 <= other.get("left", 0) - item.get("right", 0) <= 24
    ]
    candidates = [item["text"]]
    for left in left_neighbors:
        candidates.append(left["text"] + item["text"])
    for right in right_neighbors:
        candidates.append(item["text"] + right["text"])
    for left in left_neighbors:
        for right in right_neighbors:
            candidates.append(left["text"] + item["text"] + right["text"])
    best = value
    for candidate in candidates:
        parsed_candidate = parse_number_token(candidate)
        if parsed_candidate and (best is None or parsed_candidate[1] < best or "." in candidate):
            best = parsed_candidate[1]
    return best


def video_view_labels(items: list[dict]) -> list[dict]:
    labels = []
    for item in items:
        if view_label_score(item["text"]) and "video" in item["text"].lower() and "view" in item["text"].lower():
            labels.append(item)
    videos = [item for item in items if item["text"].lower().strip() in {"video", "videos"}]
    views = [item for item in items if item["text"].lower().strip() in {"view", "views"}]
    for video in videos:
        for view in views:
            if abs(video["y"] - view["y"]) <= 18 and 0 <= view["left"] - video["right"] <= 80:
                labels.append(
                    {
                        "text": f'{video["text"]} {view["text"]}',
                        "x": (video["x"] + view["x"]) / 2,
                        "y": (video["y"] + view["y"]) / 2,
                        "left": min(video["left"], view["left"]),
                        "right": max(video["right"], view["right"]),
                        "top": min(video["top"], view["top"]),
                        "bottom": max(video["bottom"], view["bottom"]),
                    }
                )
    return labels


def find_tiktok_video_views_ocr(items: list[dict]) -> int | None:
    labels = video_view_labels(items)
    if not labels:
        return None
    candidates = []
    for item in items:
        value = combined_number_value(item, items)
        if value is not None:
            candidates.append((item, value))
    best = None
    best_score = None
    for label in labels:
        for item, value in candidates:
            # TikTok Studio puts the value directly under the Video views label.
            if item["y"] <= label["bottom"]:
                continue
            if item["y"] > label["bottom"] + 90:
                continue
            if item["x"] < label["left"] - 80 or item["x"] > label["right"] + 120:
                continue
            distance = abs(item["x"] - label["x"]) + abs(item["y"] - (label["bottom"] + 35)) * 2
            score = 100000 - distance
            if best_score is None or score > best_score:
                best_score = score
                best = value
    return best


def find_number_near_views_ocr(items: list[dict]) -> int | None:
    labels = [
        (item, score)
        for item in items
        for score in [view_label_score(item["text"])]
        if score is not None
    ]
    candidates = []
    for item in items:
        value = combined_number_value(item, items)
        if value is not None:
            candidates.append((item, value))
    if not candidates:
        return None
    if not labels:
        return None
    best = None
    best_score = None
    for label, label_score in labels:
        for item, value in candidates:
            dx = abs(item["x"] - label["x"])
            dy = abs(item["y"] - label["y"])
            distance = dx + dy * 2
            # Prefer numbers directly below or to the right of the label. Penalize numbers above.
            if item["y"] < label["y"] - 12:
                distance += 500
            score = label_score * 10000 - distance
            if best_score is None or score > best_score:
                best_score = score
                best = value
    return best


def write_csv(path: Path, rows: list[list]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[list]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception:
        log_error("openpyxl import failed")
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Weekly Views"
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for col in range(1, ws.max_column + 1):
        width = max(len(str(ws.cell(row=row, column=col).value or "")) for row in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 14), 48)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center")
    wb.save(path)
    return True


class AccountDialog:
    def __init__(self, root: Tk, texts: dict, account: dict | None = None):
        self.texts = texts
        self.result = None
        self.window = Toplevel(root)
        self.window.title(self.t("account_settings"))
        self.window.transient(root)
        self.window.grab_set()
        self.window.resizable(False, False)
        self.dashboard_url = StringVar(value=(account or {}).get("dashboard_url", ""))
        self.public_url = StringVar(value=(account or {}).get("public_url", ""))

        frame = ttk.Frame(self.window, padding=16)
        frame.pack(fill=BOTH, expand=True)
        ttk.Label(frame, text=self.t("dashboard_url")).grid(row=0, column=0, sticky="w", pady=6)
        ttk.Entry(frame, textvariable=self.dashboard_url, width=62).grid(row=0, column=1, pady=6)
        ttk.Label(frame, text=self.t("public_url")).grid(row=1, column=0, sticky="w", pady=6)
        ttk.Entry(frame, textvariable=self.public_url, width=62).grid(row=1, column=1, pady=6)
        ttk.Label(frame, text=self.t("auto_detect"), foreground="#666", wraplength=620).grid(
            row=2, column=0, columnspan=2, sticky="w", pady=(12, 4)
        )
        buttons = ttk.Frame(frame)
        buttons.grid(row=3, column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(buttons, text=self.t("cancel"), command=self.window.destroy).pack(side=LEFT, padx=(0, 8))
        ttk.Button(buttons, text=self.t("save"), command=self.save).pack(side=LEFT)
        self.window.wait_window()

    def t(self, key: str) -> str:
        return self.texts.get(key, TEXTS["en"].get(key, key))

    def save(self) -> None:
        dashboard_url = self.dashboard_url.get().strip()
        public_url = self.public_url.get().strip()
        if not dashboard_url and not public_url:
            messagebox.showwarning(self.t("account_settings"), self.t("missing_url"))
            return
        self.result = infer_account(dashboard_url, public_url)
        self.window.destroy()


class SocialViewsApp:
    def __init__(self, root: Tk):
        self.root = root
        self.root.title(APP_NAME)
        self.root.geometry("900x455")
        self.root.minsize(780, 420)
        self.accounts = load_accounts()
        self.state = load_state()
        self.date_range = StringVar(value=self.state.get("date_range", default_date_range()))
        self.language = StringVar(value=LANGUAGES.get(self.state.get("language", "en"), "English"))
        self.status = StringVar(value=self.t("ready"))
        self.view_vars: dict[str, StringVar] = {}
        self.build_ui()
        self.date_range.trace_add("write", lambda *_: self.auto_save())
        self.language.trace_add("write", lambda *_: self.auto_save())

    def lang(self) -> str:
        language_name = self.language.get()
        for code, name in LANGUAGES.items():
            if language_name in {code, name}:
                return code
        return "en"

    def t(self, key: str, **kwargs) -> str:
        text = TEXTS.get(self.lang(), TEXTS["en"]).get(key, TEXTS["en"].get(key, key))
        return text.format(**kwargs) if kwargs else text

    def key_for(self, item: dict) -> str:
        return f"{item.get('platform','')}::{item.get('account','')}"

    def build_ui(self) -> None:
        self.build_menu()
        if hasattr(self, "main"):
            self.main.destroy()
        self.main = ttk.Frame(self.root, padding=18)
        self.main.pack(fill=BOTH, expand=True)

        ttk.Label(self.main, text=self.t("title"), font=("Segoe UI", 18, "bold")).pack(anchor="w")

        top = ttk.Frame(self.main)
        top.pack(fill="x", pady=(16, 20))
        ttk.Label(top, text=self.t("date_range")).pack(side=LEFT)
        ttk.Entry(top, textvariable=self.date_range, width=24).pack(side=LEFT, padx=(8, 10))
        ttk.Button(top, text=self.t("reset_dates"), command=self.reset_dates).pack(side=LEFT)
        ttk.Label(top, text=self.t("language")).pack(side=LEFT, padx=(16, 4))
        language_box = ttk.Combobox(
            top,
            textvariable=self.language,
            values=list(LANGUAGES.values()),
            width=10,
            state="readonly",
        )
        language_box.pack(side=LEFT)
        language_box.bind("<<ComboboxSelected>>", lambda _event: self.change_language())
        ttk.Button(top, text=self.t("open_all_dashboards"), command=self.open_all_dashboards).pack(
            side=LEFT, padx=(4, 0)
        )
        ttk.Button(top, text=self.t("open_all_public"), command=self.open_all_public_pages).pack(
            side=LEFT, padx=(8, 0)
        )
        ttk.Button(top, text=self.t("add_account"), command=self.add_account).pack(side=RIGHT)

        table = ttk.Frame(self.main)
        table.pack(fill=BOTH, expand=True)
        headers = [self.t("platform"), self.t("account"), self.t("views"), self.t("action")]
        widths = [14, 28, 18, 42]
        for col, (header, width) in enumerate(zip(headers, widths)):
            ttk.Label(table, text=header, font=("Segoe UI", 10, "bold")).grid(
                row=0, column=col, sticky="w", padx=4, pady=(0, 8)
            )
            table.columnconfigure(col, weight=1 if col in (1, 2, 3) else 0, minsize=width * 8)

        self.view_vars = {}
        for row_index, item in enumerate(self.accounts, start=1):
            key = self.key_for(item)
            saved = self.state.get("rows", {}).get(key, {})
            value = StringVar(value=str(saved.get("views", "")))
            self.view_vars[key] = value
            value.trace_add("write", lambda *_: self.auto_save())
            ttk.Label(table, text=item.get("platform", "")).grid(row=row_index, column=0, sticky="w", padx=4, pady=7)
            ttk.Label(table, text=item.get("account", "")).grid(row=row_index, column=1, sticky="w", padx=4, pady=7)
            ttk.Entry(table, textvariable=value, width=18).grid(row=row_index, column=2, sticky="ew", padx=4, pady=7)
            actions = ttk.Frame(table)
            actions.grid(row=row_index, column=3, sticky="w", padx=4, pady=7)
            ttk.Button(actions, text=self.t("open_dashboard"), command=lambda item=item: self.open_dashboard(item)).pack(side=LEFT)
            ttk.Button(actions, text=self.t("open_public"), command=lambda item=item: self.open_public_page(item)).pack(side=LEFT, padx=(8, 0))
            ttk.Button(actions, text=self.t("read_views"), command=lambda item=item: self.read_views(item)).pack(side=LEFT, padx=(8, 0))
            ttk.Button(actions, text=self.t("edit_account"), command=lambda item=item: self.edit_account(item)).pack(side=LEFT, padx=(8, 0))
            ttk.Button(actions, text=self.t("delete_account"), command=lambda item=item: self.delete_account(item)).pack(side=LEFT, padx=(8, 0))

        bottom = ttk.Frame(self.main)
        bottom.pack(fill="x", pady=(12, 0))
        ttk.Label(bottom, textvariable=self.status, foreground="#555").pack(side=LEFT)
        ttk.Button(bottom, text=self.t("export_report"), command=self.export_report).pack(side=RIGHT)

    def build_menu(self) -> None:
        menubar = Menu(self.root)
        advanced = Menu(menubar, tearoff=0)
        advanced.add_command(label=self.t("add_account"), command=self.add_account)
        advanced.add_command(label=self.t("open_all_dashboards"), command=self.open_all_dashboards)
        advanced.add_command(label=self.t("open_all_public"), command=self.open_all_public_pages)
        advanced.add_command(label=self.t("clear_values"), command=self.clear_values)
        advanced.add_command(label=self.t("open_config"), command=lambda: os.startfile(APP_DIR))
        menubar.add_cascade(label=self.t("advanced"), menu=advanced)
        self.root.config(menu=menubar)

    def current_state(self) -> dict:
        rows = {}
        for item in self.accounts:
            key = self.key_for(item)
            current = ""
            if key in self.view_vars:
                current = self.view_vars[key].get().strip()
            rows[key] = {
                "platform": item.get("platform", ""),
                "account": item.get("account", ""),
                "views": current,
            }
        return {"date_range": self.date_range.get().strip(), "language": self.lang(), "rows": rows}

    def auto_save(self) -> None:
        try:
            save_accounts(self.accounts)
            save_state(self.current_state())
        except Exception:
            log_error("Auto save failed")

    def change_language(self) -> None:
        self.auto_save()
        self.state = self.current_state()
        self.status.set(self.t("ready"))
        self.build_ui()

    def add_account(self) -> None:
        dialog = AccountDialog(self.root, TEXTS.get(self.lang(), TEXTS["en"]))
        if dialog.result:
            self.accounts.append(dialog.result)
            self.auto_save()
            self.state = self.current_state()
            self.build_ui()
            self.status.set(self.t("ready"))

    def edit_account(self, item: dict) -> None:
        index = self.accounts.index(item)
        dialog = AccountDialog(self.root, TEXTS.get(self.lang(), TEXTS["en"]), self.accounts[index])
        if dialog.result:
            old_key = self.key_for(self.accounts[index])
            self.accounts[index] = dialog.result
            new_key = self.key_for(dialog.result)
            if old_key in self.state.get("rows", {}) and old_key != new_key:
                self.state["rows"][new_key] = self.state["rows"].pop(old_key)
            self.auto_save()
            self.state = self.current_state()
            self.build_ui()

    def delete_account(self, item: dict) -> None:
        if not messagebox.askyesno(self.t("delete_account"), f"{item.get('platform')} / {item.get('account')}?"):
            return
        key = self.key_for(item)
        self.accounts.remove(item)
        self.state.get("rows", {}).pop(key, None)
        self.auto_save()
        self.state = self.current_state()
        self.build_ui()

    def reset_dates(self) -> None:
        self.date_range.set(default_date_range())
        self.status.set(self.t("date_reset"))

    def open_dashboard(self, item: dict) -> None:
        open_url(item.get("dashboard_url", ""))
        self.status.set(self.t("dashboard_opened", platform=item.get("platform", "")))

    def open_public_page(self, item: dict) -> None:
        if not item.get("public_url", ""):
            self.status.set(self.t("public_missing", platform=item.get("platform", "")))
            return
        open_url(item.get("public_url", ""))
        self.status.set(self.t("public_opened", platform=item.get("platform", "")))

    def open_all_dashboards(self) -> None:
        for item in self.accounts:
            open_url(item.get("dashboard_url", ""))
        self.status.set(self.t("all_dashboards_opened"))

    def open_all_public_pages(self) -> None:
        for item in self.accounts:
            if item.get("public_url", ""):
                open_url(item.get("public_url", ""))
        self.status.set(self.t("all_public_opened"))

    def read_views(self, item: dict) -> None:
        messagebox.showinfo(
            self.t("read_prompt_title"),
            self.t("read_prompt", platform=item.get("platform", "")),
        )
        self.status.set(self.t("reading_wait", platform=item.get("platform", "")))
        self.root.iconify()
        self.root.after(5000, lambda: self.finish_read_views(item))

    def finish_read_views(self, item: dict) -> None:
        try:
            key = self.key_for(item)
            self.status.set(self.t("reading", platform=item.get("platform", "")))
            sentinel = f"__EMPTY_CLIPBOARD_{int(time.time())}__"
            try:
                self.root.clipboard_clear()
                self.root.clipboard_append(sentinel)
                self.root.update()
            except Exception:
                pass
            copy_active_window_text()
            try:
                page_text = self.root.clipboard_get()
            except Exception:
                page_text = ""
            if page_text == sentinel:
                page_text = ""
            views = parse_views_from_text(page_text)
            if views is None:
                ocr_items = read_screen_ocr_items()
                if item.get("platform", "").strip().lower() == "tiktok":
                    views = find_tiktok_video_views_ocr(ocr_items)
                if views is None:
                    views = find_number_near_views_ocr(ocr_items)
                if views is None:
                    views = parse_views_from_text("\n".join(item["text"] for item in ocr_items))
            if views is None:
                self.status.set(self.t("read_failed"))
                self.root.deiconify()
                self.root.lift()
                messagebox.showinfo(self.t("read_prompt_title"), self.t("read_failed"))
                return
            self.root.deiconify()
            self.root.lift()
            self.view_vars[key].set(str(views))
            self.auto_save()
            self.status.set(self.t("detected", platform=item.get("platform", ""), views=views))
        except Exception:
            log_error("Read visible views failed")
            self.root.deiconify()
            self.root.lift()
            self.status.set(self.t("read_failed"))

    def clear_values(self) -> None:
        if not messagebox.askyesno(self.t("clear_title"), self.t("clear_confirm")):
            return
        for key in self.view_vars:
            self.view_vars[key].set("")
        self.auto_save()
        self.status.set(self.t("cleared"))

    def export_rows(self) -> list[list]:
        rows = [[self.t("col_platform"), self.t("col_account"), self.t("col_views"), self.t("col_date"), self.t("col_raw")]]
        range_text = self.date_range.get().strip()
        for item in self.accounts:
            key = self.key_for(item)
            raw = self.view_vars.get(key, StringVar(value="")).get().strip()
            rows.append(
                [
                    item.get("platform", ""),
                    item.get("account", ""),
                    safe_int(raw) if raw else 0,
                    range_text,
                    raw,
                ]
            )
        return rows

    def export_report(self) -> None:
        try:
            self.auto_save()
            start, end = parse_date_range(self.date_range.get())
            filename = f"social_weekly_views_{start}_to_{end}"
            rows = self.export_rows()
            desktop_xlsx = DESKTOP / f"{filename}.xlsx"
            desktop_csv = DESKTOP / f"{filename}.csv"
            backup_xlsx = EXPORT_DIR / f"{filename}.xlsx"
            backup_csv = EXPORT_DIR / f"{filename}.csv"
            xlsx_ok = write_xlsx(desktop_xlsx, rows)
            if xlsx_ok:
                write_xlsx(backup_xlsx, rows)
            write_csv(desktop_csv, rows)
            write_csv(backup_csv, rows)
            self.status.set(self.t("exported"))
            messagebox.showinfo(
                self.t("export_title"),
                self.t("export_success", path=desktop_xlsx if xlsx_ok else desktop_csv),
            )
        except Exception:
            log_error("Export report failed")
            self.status.set(self.t("export_failed"))
            messagebox.showwarning(self.t("export_title"), self.t("export_failed"))


def main() -> None:
    try:
        ensure_dirs()
        root = Tk()
        try:
            style = ttk.Style(root)
            if "vista" in style.theme_names():
                style.theme_use("vista")
        except Exception:
            pass
        SocialViewsApp(root)
        root.mainloop()
    except Exception:
        log_error("Application startup failed")
        sys.exit(1)


if __name__ == "__main__":
    main()
